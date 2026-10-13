import sys
from pathlib import Path

import pytest
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

HEADERS = ["テストケース番号 *\n(TC ID)", "分類 *\n(Category)", "テスト項目 *\n(Title)", "前提条件\n(Pre-condition)",
           "手順 *\n(Steps)", "テストデータ\n(Test Data)", "期待値 *\n(Expected Result)", "優先度 *\n(Priority)",
           "テスト結果 (Test Result) – Chrome", None, None, "最終結果\n(Final Result)", "実施者 *\n(Tester / PIC)",
           "実施日\n(Date)", "不具合ID *\n(Bug ID)", "備考\n(Note)"]

CASES = [
    # id, category, title, precondition, steps, data, expected, priority, run1, run2, bug, note
    ("Shop_TC001", "1. 一覧", "List shows products", "Logged in", "1. Open the list.", None,
     "1. 10 products are shown.\n2. Sorted by name.", "High", "Passed", None, None, None),
    ("Shop_TC002", "1. 一覧", "Upload too large", "On the list", "1. Upload FILE-NG-01.", "FILE-NG-01 (>5MB)",
     "1. Toast is shown.", "Medium", "Failed", "Passed", "BUG-12", "Bug: https://tracker/BUG-12"),
    ("Shop_TC004", "2. 詳細", "Detail matches design", "Shop_TC001 done", "1. Open detail.\n   then scroll",
     "vd 'abc'", "1. Header hiển thị đúng thiết kế.", "Low", "Failed", None, None, None),
    ("Shop_TC005", "2. 詳細", "No steps", None, None, "Q9", "1. Something.", "Urgent", None, None, None, None),
]


def build_workbook(path: Path, cases=CASES) -> Path:
    workbook = Workbook()
    cover = workbook.active
    cover.title = "表紙_Cover"
    cover["B9"] = "結合テスト仕様書"

    description = workbook.create_sheet("項目説明_Item Description")
    description.append(["No.", "項目\n(Item)", "説明 (JP)"])
    description.append([1, "テストケース番号\n(TC ID)", "固定ID"])

    sheet = workbook.create_sheet("001-SHP_Shop")
    sheet["A1"] = "結合テスト仕様書兼結果書 (IT Specification & Result)"
    sheet["A2"], sheet["C2"], sheet["F2"] = "画面・機能名\n(Screen / Function Name)", "Shop 画面 (Shop)", "テスト結果集計 (Result Statistics)"
    sheet["A3"], sheet["C3"] = "モジュール\n(Module)", "Shop"
    sheet["A4"], sheet["C4"] = "テスト環境\n(Environment)", "Staging / Chrome"
    for column, value in enumerate(HEADERS, start=1):
        sheet.cell(row=7, column=column, value=value)
    sheet["I8"], sheet["J8"], sheet["K8"] = "1回目 *\n(1st)", "2回目\n(2nd)", "3回目\n(3rd)"
    for index, case in enumerate(cases):
        row = 9 + index
        (case_id, category, title, pre, steps, data, expected, priority, run1, run2, bug, note) = case
        for column, value in zip("ABCDEFGH", (case_id, category, title, pre, steps, data, expected, priority)):
            sheet[f"{column}{row}"] = value
        sheet[f"I{row}"], sheet[f"J{row}"] = run1, run2
        sheet[f"L{row}"] = f'=IF($A{row}="","",IF(K{row}<>"",K{row},IF(J{row}<>"",J{row},IF(I{row}<>"",I{row},"Todo"))))'
        sheet[f"O{row}"], sheet[f"P{row}"] = bug, note

    data = workbook.create_sheet("TestData")
    data["A1"] = "テストデータ (Test Data)"
    data.append(["データID\n(Data ID)", "種別\n(Type)", "内容\n(Content)", "言語\n(Language)", "使用TC\n(Used in)", "備考\n(Note)"])
    data.append(["Q1", "質問\n(Question)", "こんにちは", "JA", "Shop_TC001", "Câu hỏi chuẩn"])
    data.append(["FILE-NG-01", "ファイル\n(File)", "PDF >5MB", "-", "Shop_TC002", "Cần chuẩn bị"])

    workbook.save(path)
    return path


@pytest.fixture()
def workbook_path(tmp_path) -> Path:
    return build_workbook(tmp_path / "cases.xlsx")
