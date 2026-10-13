"""AI information on imported test cases: execution notes, edits to existing cases, new cases, the tester's
decisions, and writing accepted content into the workbook."""

import json
import subprocess
import sys
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook
from openpyxl.styles import Border, Side

from conftest import AUTOMATION_COLUMN, CASES, build_workbook, needs_pty, run_cli_tty
from wtt import plan as P
from wtt import proposals as R
from wtt.review import render_review
from wtt.workbook import cell_text, read_workbook, write_case_changes

CLI = Path(__file__).resolve().parents[1] / "wtt_import.py"
SHEET = "001-SHP_Shop"


def build(path, cases=CASES):
    raw = read_workbook(build_workbook(path, cases))
    plan = P.build_plan(raw, raw.case_sheets[0], imported_at="2026-10-13T10:00:00+07:00")
    P.apply_workbook_marks(plan)
    P.analyse(plan)
    return plan


@pytest.fixture()
def plan(tmp_path):
    return build(tmp_path / "w.xlsx")


def by_id(plan):
    return {case["id"]: case for case in plan["test_cases"]}


def codes(item):
    return {issue["code"] for issue in item.get("issues") or []}


def edit(field="expected", value=None, **extra):
    proposal = {"field": field, "value": value if value is not None else ["12 products are shown."],
                "reason": "The count is in the spec.", "source": "document", "evidence": "Spec 3.2", "status": "proposed"}
    proposal.update(extra)
    return proposal


def new_case(case_id="Shop_TC006", **extra):
    proposal = {"id": case_id, "category": "2. 詳細", "title": "Back button returns to the list",
                "preconditions": ["On a product detail"], "steps": ["Click Back."],
                "expected": ["The list is shown.", "The scroll position is kept."], "priority": "Medium",
                "automation": "auto", "review": "Plain navigation.", "reason": "No case covers Back.",
                "source": "site", "evidence": "explore /shop/1", "status": "proposed"}
    proposal.update(extra)
    return proposal


# ---------- checks ----------

def test_valid_edit_has_no_errors(plan):
    by_id(plan)["Shop_TC001"]["proposals"] = [edit()]
    P.analyse(plan)
    assert codes(by_id(plan)["Shop_TC001"]["proposals"][0]) == set()
    assert plan["summary"]["proposals"] == {"edit": {"proposed": 1}, "new_case": {}}


@pytest.mark.parametrize("change, code", [
    ({"reason": None}, "proposal_reason"),
    ({"source": "guess"}, "proposal_source"),
    ({"evidence": None}, "proposal_evidence"),
    ({"field": "bug_id"}, "proposal_field"),
    ({"value": []}, "proposal_value"),
    ({"field": "priority", "value": "Urgent"}, "proposal_priority"),
    ({"status": "maybe"}, "proposal_status"),
    ({"value": ["10 products are shown.", "Sorted by name."]}, "proposal_no_change"),
    ({"source": "site", "evidence": "explore /"}, "proposal_observed"),
])
def test_edit_checks(plan, change, code):
    proposal = edit()
    proposal.update(change)
    by_id(plan)["Shop_TC001"]["proposals"] = [proposal]
    P.analyse(plan)
    assert code in codes(proposal)


def test_inference_needs_no_evidence(plan):
    proposal = edit(source="inference", evidence=None)
    by_id(plan)["Shop_TC001"]["proposals"] = [proposal]
    P.analyse(plan)
    assert codes(proposal) == set()


@pytest.mark.parametrize("change, code", [
    ({"id": "SHOP-6"}, "new_id_format"),
    ({"id": "Cart_TC006"}, "new_id_function"),
    ({"id": "Shop_TC002"}, "new_id_exists"),
    ({"id": "Shop_TC003"}, "new_id_reused"),
    ({"steps": []}, "new_missing_steps"),
    ({"expected": None}, "new_missing_expected"),
    ({"title": ""}, "new_missing_title"),
    ({"priority": "P0"}, "new_priority"),
    ({"automation": "yes"}, "new_automation"),
    ({"category": "9. その他"}, "new_category"),
    ({"execution": "click it"}, "execution_format"),
])
def test_new_case_checks(plan, change, code):
    proposal = new_case(**change)
    plan["proposed_cases"] = [proposal]
    P.analyse(plan)
    assert code in codes(proposal)


def test_duplicate_new_ids(plan):
    plan["proposed_cases"] = [new_case(), new_case()]
    P.analyse(plan)
    assert all("new_id_duplicate" in codes(proposal) for proposal in plan["proposed_cases"])


def test_next_free_id_skips_gaps_and_proposals(plan):
    assert R.next_free_id(plan) == "Shop_TC006"
    plan["proposed_cases"] = [new_case()]
    assert R.next_free_id(plan) == "Shop_TC007"


def test_execution_checks_and_hash(plan):
    case = by_id(plan)["Shop_TC001"]
    before = P.content_hash(plan)
    case["execution"] = {"navigation": "/shop", "locators": {"list": "getByRole('list')"}, "checks": ["count == 10"]}
    P.analyse(plan)
    assert codes(case) == set()
    assert P.content_hash(plan) != before, "execution is approved with the plan"
    assert plan["summary"]["with_execution"] == 1
    case["execution"] = {"navigation": "/shop", "selector": "#list"}
    P.analyse(plan)
    assert "execution_unknown_key" in codes(case)
    case["execution"] = ["open /shop"]
    P.analyse(plan)
    assert "execution_format" in codes(case)


def test_review_page_shows_every_execution_note(plan):
    case = by_id(plan)["Shop_TC001"]
    case["execution"] = {"navigation": "/shop", "locators": {"list": "getByRole('list')"},
                         "checks": ["count == 10", "sorted by name"], "selector": "#list"}
    P.analyse(plan)
    page = render_review(plan)
    assert "## Thông tin chạy do AI ghi (duyệt cùng plan)" in page
    assert ("| Shop_TC001 | /shop | list: getByRole('list') |  |  | - count == 10<br>- sorted by name |  | "
            "selector: #list |") in page
    assert "## Thông tin chạy" not in render_review(build_plain(plan))


def build_plain(plan):
    for case in plan["test_cases"]:
        case.pop("execution", None)
    P.analyse(plan)
    return plan


# ---------- decisions ----------

def test_decide_accept_and_reject(plan):
    by_id(plan)["Shop_TC001"]["proposals"] = [edit(), edit(field="title", value="List shows 12 products")]
    plan["proposed_cases"] = [new_case()]
    P.analyse(plan)
    assert R.decide(plan, ["Shop_TC001#2", "Shop_TC006"], accept=True, by="NCM", at="2026-10-13T12:00:00+07:00") == [
        "Shop_TC001#2", "Shop_TC006"]
    assert R.decide(plan, ["all"], accept=False, by="NCM", note="not in spec") == ["Shop_TC001#1"]
    first, second = by_id(plan)["Shop_TC001"]["proposals"]
    assert first["status"] == "rejected" and first["decided"]["note"] == "not in spec"
    assert second["status"] == "accepted" and second["decided"] == {"by": "NCM", "at": "2026-10-13T12:00:00+07:00"}
    with pytest.raises(ValueError, match="Không còn đề xuất"):
        R.decide(plan, ["all"], accept=True, by="NCM")


def test_decide_refuses_errors_atomically(plan):
    by_id(plan)["Shop_TC001"]["proposals"] = [edit()]
    plan["proposed_cases"] = [new_case(id="Shop_TC003")]
    P.analyse(plan)
    with pytest.raises(ValueError, match="Shop_TC003: còn lỗi"):
        R.decide(plan, ["all"], accept=True, by="NCM")
    assert by_id(plan)["Shop_TC001"]["proposals"][0]["status"] == "proposed", "nothing decided when one is refused"
    R.decide(plan, ["Shop_TC003"], accept=False, by="NCM")


@pytest.mark.parametrize("handle, message", [
    ("Shop_TC001", "đề xuất sửa case dùng mã Shop_TC001#<n>"),
    ("Shop_TC001#9", "Không có đề xuất"),
    ("nope", "không phải mã đề xuất"),
])
def test_decide_bad_handles(plan, handle, message):
    with pytest.raises(ValueError, match=message):
        R.decide(plan, [handle], accept=True, by="NCM")


def test_decide_needs_a_name_and_keeps_written(plan):
    by_id(plan)["Shop_TC001"]["proposals"] = [edit(status="accepted", written={"file": "x.xlsx", "at": "t"})]
    P.analyse(plan)
    with pytest.raises(ValueError, match="--by"):
        R.decide(plan, ["Shop_TC001#1"], accept=False, by=" ")
    with pytest.raises(ValueError, match="đã ghi vào workbook"):
        R.decide(plan, ["Shop_TC001#1"], accept=False, by="NCM")


def test_open_proposals_block_approval(plan):
    for case in plan["test_cases"]:
        case["automation"], case["review"] = ("need_info", "no steps") if case["id"] == "Shop_TC005" else ("auto", None)
    by_id(plan)["Shop_TC002"]["proposals"] = [edit(value=["Toast 'File quá lớn' is shown."])]
    plan["proposed_cases"] = [new_case()]
    P.analyse(plan)
    blockers = P.approval_blockers(plan)
    assert any(item.startswith("Shop_TC002#1 (đề xuất sửa): tester chưa quyết định") for item in blockers)
    assert any(item.startswith("Shop_TC006 (case mới)") for item in blockers)

    R.decide(plan, ["Shop_TC002#1"], accept=True, by="NCM")
    R.decide(plan, ["Shop_TC006"], accept=False, by="NCM")
    assert P.approval_blockers(plan) == [
        "Shop_TC002#1 (đề xuất sửa): đã chấp nhận nhưng chưa ghi vào workbook (write-changes, rồi import lại)."]
    by_id(plan)["Shop_TC002"]["proposals"][0]["written"] = {"file": "w_updated.xlsx", "at": "t"}
    assert P.approval_blockers(plan) == []


def test_review_page_lists_proposals(plan):
    by_id(plan)["Shop_TC001"]["proposals"] = [edit()]
    plan["proposed_cases"] = [new_case(id="Shop_TC003")]
    P.analyse(plan)
    page = render_review(plan)
    assert "## Đề xuất của AI (tester quyết định)" in page
    assert ("| Shop_TC001#1 | Expected Result | 1. 10 products are shown.<br>2. Sorted by name. | "
            "1. 12 products are shown. | document: Spec 3.2 |") in page
    assert "| Shop_TC003 | 2. 詳細 | Back button returns to the list |" in page
    assert "mã đã bỏ trong dãy số" in page


# ---------- writing ----------

@pytest.mark.parametrize("column, value, text", [
    ("steps", ["Open the list."], "1. Open the list."),
    ("expected", ["a", "b"], "1. a\n2. b"),
    ("precondition", ["Logged in"], "Logged in"),
    ("precondition", ["Logged in", "Cart empty"], "1. Logged in\n2. Cart empty"),
    ("priority", "medium", "Medium"),
    ("test_data", "FILE-NG-01", "FILE-NG-01"),
    ("title", None, None),
])
def test_cell_text(column, value, text):
    assert cell_text(column, value) == text


def styled_workbook(path):
    build_workbook(path)
    workbook = load_workbook(path)
    sheet = workbook[SHEET]
    thin = Side(style="thin")
    for cell in sheet[12]:
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    workbook.save(path)
    return path


def test_write_case_changes(tmp_path):
    source = styled_workbook(tmp_path / "w.xlsx")
    out = tmp_path / "out.xlsx"
    added = {"id": "Shop_TC006", "category": "2. 詳細", "title": "Back", "precondition": ["On detail"],
             "steps": ["Click Back."], "test_data": None, "expected": ["List shown."], "priority": "medium",
             "automation_mark": "Automated"}
    result = write_case_changes(source, out, SHEET,
                                [("Shop_TC001", "expected", ["12 products are shown."]), ("Shop_TC002", "priority", "high")],
                                [added], {"Shop_TC001": "[AI] sửa Expected", "Shop_TC006": "[AI] case mới"})
    assert result == {"edited": [["Shop_TC001", "expected"], ["Shop_TC002", "priority"]],
                      "added": [["Shop_TC006", 13]], "beyond_template_rows": ["Shop_TC006"]}

    sheet = load_workbook(out)[SHEET]
    assert sheet["G9"].value == "1. 12 products are shown."
    assert sheet["H10"].value == "High"
    assert sheet["Q9"].value == "[AI] sửa Expected"
    assert sheet["Q10"].value == "Bug: https://tracker/BUG-12", "notes go only to the cases they name"
    assert [sheet[f"{column}13"].value for column in "ABCDEGH"] == [
        "Shop_TC006", "2. 詳細", "Back", "On detail", "1. Click Back.", "1. List shown.", "Medium"]
    assert sheet[f"{AUTOMATION_COLUMN}13"].value == "Automated"
    assert sheet["Q13"].value == "[AI] case mới"
    assert sheet["L13"].value == '=IF($A13="","",IF(K13<>"",K13,IF(J13<>"",J13,IF(I13<>"",I13,"Todo"))))'
    assert sheet["A13"].border.left.style == "thin", "style copied from the last case row"

    raw = read_workbook(out)
    assert [case.values["id"] for case in raw.case_sheets[0].cases][-1] == "Shop_TC006"


def test_write_case_changes_appends_to_an_existing_note(tmp_path):
    out = tmp_path / "out.xlsx"
    write_case_changes(build_workbook(tmp_path / "w.xlsx"), out, SHEET, [("Shop_TC002", "title", "Upload > 5MB")], [],
                       {"Shop_TC002": "[AI] sửa Title"})
    assert load_workbook(out)[SHEET]["Q10"].value == "Bug: https://tracker/BUG-12\n[AI] sửa Title"


def test_write_case_changes_writes_nothing_on_a_problem(tmp_path):
    source = build_workbook(tmp_path / "w.xlsx")
    out = tmp_path / "out.xlsx"
    with pytest.raises(ValueError) as error:
        write_case_changes(source, out, SHEET, [("Shop_TC009", "title", "x")],
                           [{"id": "Shop_TC001", "title": "dup"}], {})
    assert "Shop_TC009: không có trong sheet" in str(error.value)
    assert "Shop_TC001: mã đã có trong sheet" in str(error.value)
    assert not out.exists()


# ---------- re-import ----------

def test_carry_over_keeps_open_proposals(tmp_path):
    old = build(tmp_path / "w.xlsx")
    by_id(old)["Shop_TC002"]["proposals"] = [edit(status="rejected"), edit(status="accepted", written={"file": "x", "at": "t"})]
    old["proposed_cases"] = [new_case(), new_case(id="Shop_TC007", status="rejected")]
    new = build(tmp_path / "w2.xlsx")
    P.merge_review(new, old)
    assert R.carry_over(new, old) == []
    assert [proposal["status"] for proposal in by_id(new)["Shop_TC002"]["proposals"]] == ["rejected"]
    assert [proposal["id"] for proposal in new["proposed_cases"]] == ["Shop_TC006", "Shop_TC007"]


# ---------- CLI ----------

def run_cli(*args, cwd):
    return subprocess.run([sys.executable, str(CLI), *args], cwd=cwd, capture_output=True, text=True, timeout=60)


def seed(plan_path, mutate):
    data = yaml.safe_load(plan_path.read_text(encoding="utf-8"))
    mutate(data)
    plan_path.write_text(yaml.safe_dump(data, allow_unicode=True, sort_keys=False), encoding="utf-8")


@needs_pty
def test_cli_proposals_end_to_end(tmp_path):
    work = tmp_path / "work"
    work.mkdir()
    workbook = build_workbook(work / "cases.xlsx")
    imported = json.loads(run_cli("import", str(workbook), cwd=work).stdout)[0]
    assert imported["next_free_id"] == "Shop_TC006"
    plan_path = work / imported["plan"]

    def add_ai_information(data):
        cases = {case["id"]: case for case in data["test_cases"]}
        cases["Shop_TC004"]["proposals"] = [edit(value=["Header shows 'Shop detail'."], source="site", evidence="explore /shop/1")]
        cases["Shop_TC002"]["execution"] = {"data": "Generate a 6 MB PDF in the fixture."}
        data["proposed_cases"] = [new_case(), new_case(id="Shop_TC007", title="Share link", automation="manual",
                                                       review="Needs a second account.")]
    seed(plan_path, add_ai_information)
    assert run_cli("check", str(plan_path), cwd=work).returncode == 0

    nothing = run_cli("write-changes", str(plan_path), str(workbook), cwd=work)
    assert nothing.returncode == 0 and "Không có đề xuất" in nothing.stdout

    decided = run_cli_tty("decide", str(plan_path), "Shop_TC004#1", "Shop_TC006", "--accept", "--by", "NCM",
                          cwd=work, answer="2")
    assert decided.returncode == 0, decided.output
    assert json.loads(decided.output[decided.output.index("{"):])["open"] == [
        "Shop_TC004#1 (đề xuất sửa): đã chấp nhận nhưng chưa ghi vào workbook (write-changes, rồi import lại).",
        "Shop_TC006 (case mới): đã chấp nhận nhưng chưa ghi vào workbook (write-changes, rồi import lại).",
        "Shop_TC007 (case mới): tester chưa quyết định (decide --accept / --reject).",
    ]
    assert run_cli_tty("decide", str(plan_path), "Shop_TC007", "--reject", "--by", "NCM", cwd=work, answer="1").returncode == 0

    written = run_cli("write-changes", str(plan_path), str(workbook), cwd=work)
    assert written.returncode == 0, written.stderr
    result = json.loads(written.stdout)
    assert result["workbook"].endswith("cases_updated.xlsx")
    assert result["edited"] == [["Shop_TC004", "expected"]] and result["added"] == [["Shop_TC006", 13]]
    assert result["next"].startswith("import ") and result["next"].endswith("cases_updated.xlsx")
    sheet = load_workbook(work / "cases_updated.xlsx")[SHEET]
    assert sheet["G11"].value == "1. Header shows 'Shop detail'."
    assert sheet["Q11"].value.endswith("NCM duyệt sửa: Expected Result (nguồn: site).")
    assert sheet["Q13"].value.endswith("Case mới do AI đề xuất (nguồn: site), NCM duyệt.")
    assert load_workbook(workbook)[SHEET]["G11"].value == "1. Header hiển thị đúng thiết kế.", "input left as it was"

    again = run_cli("write-changes", str(plan_path), str(workbook), cwd=work)
    assert "Không có đề xuất" in again.stdout, "written proposals are not written twice"

    reimported = json.loads(run_cli("import", str(work / "cases_updated.xlsx"), cwd=work).stdout)[0]
    assert reimported["changed_since_last_import"] == ["Shop_TC004"]
    assert reimported["added_from_proposals"] == ["Shop_TC006"]
    assert reimported["summary"]["total"] == 5
    assert reimported["summary"]["proposals"] == {"edit": {}, "new_case": {"rejected": 1}}
    assert reimported["next_free_id"] == "Shop_TC008"
    data = yaml.safe_load(plan_path.read_text(encoding="utf-8"))
    cases = {case["id"]: case for case in data["test_cases"]}
    assert (cases["Shop_TC006"]["automation"], cases["Shop_TC006"]["review"]) == ("auto", "Plain navigation.")
    assert cases["Shop_TC006"]["workbook_automation"] == "Automated"
    assert cases["Shop_TC004"]["automation"] is None, "edited content is reviewed again"
    assert cases["Shop_TC004"]["expected"] == ["Header shows 'Shop detail'."]
    assert "vague_expected" not in codes(cases["Shop_TC004"])
    assert cases["Shop_TC002"]["execution"] == {"data": "Generate a 6 MB PDF in the fixture."}


@needs_pty
def test_cli_write_changes_refuses_a_changed_workbook(tmp_path):
    work = tmp_path / "work"
    work.mkdir()
    workbook = build_workbook(work / "cases.xlsx")
    plan_path = work / json.loads(run_cli("import", str(workbook), cwd=work).stdout)[0]["plan"]
    seed(plan_path, lambda data: data.__setitem__("proposed_cases", [new_case()]))
    assert run_cli_tty("decide", str(plan_path), "all", "--accept", "--by", "NCM", cwd=work, answer="1").returncode == 0
    edited = list(CASES)
    edited[0] = edited[0][:2] + ("List shows all products",) + edited[0][3:]
    build_workbook(workbook, edited)
    refused = run_cli("write-changes", str(plan_path), str(workbook), cwd=work)
    assert refused.returncode == 3 and "import lại" in refused.stderr


@needs_pty
def test_cli_decide_errors(tmp_path):
    work = tmp_path / "work"
    work.mkdir()
    workbook = build_workbook(work / "cases.xlsx")
    plan_path = work / json.loads(run_cli("import", str(workbook), cwd=work).stdout)[0]["plan"]
    refused = run_cli_tty("decide", str(plan_path), "all", "--accept", "--by", "NCM", cwd=work)
    assert refused.returncode == 1 and "Không còn đề xuất" in refused.output
    assert run_cli("decide", str(plan_path), "Shop_TC001#1", "--by", "NCM", cwd=work).returncode == 2
