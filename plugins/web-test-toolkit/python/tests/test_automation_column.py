"""The workbook's Automation column: read on import, kept in step with the plan, written back after approval,
and used to select what the executor runs."""

import copy
import hashlib
import json
import subprocess
import sys
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook

from conftest import AUTOMATION_COLUMN, CASES, build_workbook, needs_pty, run_cli_tty
from wtt import plan as P
from wtt.review import render_review
from wtt.workbook import case_column, read_workbook, write_automation_marks

CLI = Path(__file__).resolve().parents[1] / "wtt_import.py"


def build(path, cases=CASES, **kwargs):
    raw = read_workbook(build_workbook(path, cases, **kwargs))
    return P.build_plan(raw, raw.case_sheets[0], imported_at="2026-10-13T10:00:00+07:00")


def by_id(plan):
    return {case["id"]: case for case in plan["test_cases"]}


def codes(case):
    return {issue["code"] for issue in case["issues"]}


def with_mark(case, mark):
    return tuple(case[:12]) + (mark,)


def reimport(old_plan, path, cases):
    """What `import` does on a re-import: merge the previous review, then let changed ticks decide."""
    new_plan = build(path, cases)
    changed = P.merge_review(new_plan, old_plan)
    applied = P.apply_workbook_marks(new_plan, old_plan)
    P.analyse(new_plan)
    return new_plan, changed, applied


# ---------- reading the column ----------

def test_case_column_accepts_plain_and_bilingual_headers():
    assert case_column("Automation") == "automation_mark"
    assert case_column(" automation ") == "automation_mark"
    assert case_column("自動化\n(Automation)") == "automation_mark"
    assert case_column("テストケース番号 *\n(TC ID)") == "id"
    assert case_column("Passed") is None
    assert case_column(None) is None


def test_import_keeps_the_tick_as_imported(tmp_path):
    plan = build(tmp_path / "w.xlsx")
    cases = by_id(plan)
    assert cases["Shop_TC001"]["workbook_automation"] == "Automated"
    assert cases["Shop_TC004"]["workbook_automation"] == "Not Automated"
    assert cases["Shop_TC002"]["workbook_automation"] is None
    assert all(case["automation"] is None for case in plan["test_cases"]), "build_plan alone decides nothing"
    assert plan["summary"]["by_workbook_automation"] == {"Automated": 1, "none": 2, "Not Automated": 1}


def test_older_template_without_the_column_still_imports(tmp_path):
    plan = build(tmp_path / "old.xlsx", automation_column=False)
    assert all(case["workbook_automation"] is None for case in plan["test_cases"])
    assert by_id(plan)["Shop_TC002"]["history"]["bug_ids"] == ["BUG-12"]


@pytest.mark.parametrize("text, expected", [
    ("Automated", "Automated"), ("automated", "Automated"), ("Not Automated", "Not Automated"),
    ("not  automated", "Not Automated"), ("Not-Automated", "Not Automated"), (None, None), ("", None), ("Yes", "Yes"),
])
def test_normalize_mark(text, expected):
    assert P.normalize_mark(text) == expected


@pytest.mark.parametrize("decision, mark", [
    ("auto", "Automated"), ("partial", "Automated"), ("manual", "Not Automated"), ("need_info", "Not Automated"), (None, None),
])
def test_decision_mark(decision, mark):
    assert P.decision_mark(decision) == mark


# ---------- tick -> decision ----------

def test_first_import_takes_decisions_from_the_ticks(tmp_path):
    plan = build(tmp_path / "w.xlsx")
    assert P.apply_workbook_marks(plan) == ["Shop_TC001", "Shop_TC004"]
    cases = by_id(plan)
    assert (cases["Shop_TC001"]["automation"], cases["Shop_TC001"]["review"]) == (
        "auto", "Theo cột Automation trong workbook: Automated.")
    assert (cases["Shop_TC004"]["automation"], cases["Shop_TC004"]["review"]) == (
        "manual", "Theo cột Automation trong workbook: Not Automated.")
    assert cases["Shop_TC002"]["automation"] is None, "no tick: the AI review and the tester decide"


def test_unknown_tick_is_flagged_and_ignored(tmp_path):
    cases = list(CASES)
    cases[1] = with_mark(cases[1], "Yes")
    plan = build(tmp_path / "w.xlsx", cases)
    assert "Shop_TC002" not in P.apply_workbook_marks(plan)
    P.analyse(plan)
    assert "workbook_automation_value" in codes(by_id(plan)["Shop_TC002"])


def test_plan_decision_that_differs_from_the_tick_is_flagged(tmp_path):
    plan = build(tmp_path / "w.xlsx")
    P.apply_workbook_marks(plan)
    case = by_id(plan)["Shop_TC001"]
    case["automation"], case["review"] = "need_info", "Expected value for the sort order is missing."
    P.analyse(plan)
    message = [issue for issue in case["issues"] if issue["code"] == "workbook_automation_differs"][0]["message"]
    assert "'Automated'" in message and "'need_info'" in message and "'Not Automated'" in message
    page = render_review(plan)
    assert "| Shop_TC001 | List shows products | high | Automated | need_info | Not Automated ✏️ |" in page


def test_partial_agrees_with_automated(tmp_path):
    plan = build(tmp_path / "w.xlsx")
    P.apply_workbook_marks(plan)
    case = by_id(plan)["Shop_TC001"]
    case["automation"], case["review"] = "partial", "Sort order checked, look by screenshot."
    P.analyse(plan)
    assert "workbook_automation_differs" not in codes(case)


# ---------- re-import ----------

def first_import(path):
    plan = build(path)
    P.apply_workbook_marks(plan)
    P.analyse(plan)
    return plan


def test_unchanged_tick_does_not_undo_a_later_plan_decision(tmp_path):
    old = first_import(tmp_path / "w.xlsx")
    case = by_id(old)["Shop_TC001"]
    case["automation"], case["review"] = "need_info", "Sort order unknown."
    new, changed, applied = reimport(old, tmp_path / "w2.xlsx", CASES)
    assert changed == [] and applied == []
    assert by_id(new)["Shop_TC001"]["automation"] == "need_info"
    assert "workbook_automation_differs" in codes(by_id(new)["Shop_TC001"])


def test_changed_tick_overrides_the_plan_decision(tmp_path):
    old = first_import(tmp_path / "w.xlsx")
    cases_old = by_id(old)
    cases_old["Shop_TC002"]["automation"] = "auto"
    cases_old["Shop_TC005"]["automation"], cases_old["Shop_TC005"]["review"] = "need_info", "No steps."
    P.approve(old, "NCM", at="2026-10-13T11:00:00+07:00")

    cases = list(CASES)
    cases[1] = with_mark(cases[1], "Not Automated")   # TC002: (blank) auto -> Not Automated
    cases[2] = with_mark(cases[2], "Automated")       # TC004: Not Automated -> Automated
    cases[3] = with_mark(cases[3], "Not Automated")   # TC005: (blank) need_info -> Not Automated: agrees, kept
    new, changed, applied = reimport(old, tmp_path / "w2.xlsx", cases)

    assert changed == [] and applied == ["Shop_TC002", "Shop_TC004"]
    second, fourth, fifth = (by_id(new)[case_id] for case_id in ("Shop_TC002", "Shop_TC004", "Shop_TC005"))
    assert second["automation"] == "manual"
    assert second["review"] == "Theo cột Automation trong workbook: Not Automated. (quyết định trước: auto)"
    assert fourth["automation"] == "auto"
    assert fourth["review"].startswith("Theo cột Automation trong workbook: Automated. (quyết định trước: manual")
    assert (fifth["automation"], fifth["review"]) == ("need_info", "No steps.")
    new["approval"] = old["approval"]
    assert P.approval_state(new)[0] == "stale", "a tick changed in Excel needs a new approval"


def test_changed_content_with_the_same_tick_must_be_reviewed_again(tmp_path):
    old = first_import(tmp_path / "w.xlsx")
    cases = list(CASES)
    cases[0] = cases[0][:2] + ("List shows all products",) + cases[0][3:]
    new, changed, applied = reimport(old, tmp_path / "w2.xlsx", cases)
    assert changed == ["Shop_TC001"] and applied == []
    assert by_id(new)["Shop_TC001"]["automation"] is None


def test_new_case_with_a_tick_takes_it(tmp_path):
    old = first_import(tmp_path / "w.xlsx")
    added = ("Shop_TC006", "2. 詳細", "Back button", None, "1. Click back.", None, "1. The list is shown.", "Medium",
             None, None, None, None, "Automated")
    new, changed, applied = reimport(old, tmp_path / "w2.xlsx", list(CASES) + [added])
    assert applied == ["Shop_TC006"] and by_id(new)["Shop_TC006"]["automation"] == "auto"


# ---------- execution selection ----------

def test_execution_cases(tmp_path):
    plan = build(tmp_path / "w.xlsx")
    decisions = {"Shop_TC001": "auto", "Shop_TC002": "partial", "Shop_TC004": "auto", "Shop_TC005": "need_info"}
    for case in plan["test_cases"]:
        case["automation"] = decisions[case["id"]]
    # TC001 Passed, TC002 Failed -> Passed, TC004 Failed (retest), TC005 not run
    assert [case["id"] for case in P.execution_cases(plan)] == ["Shop_TC001", "Shop_TC002", "Shop_TC004"]
    assert [case["id"] for case in P.execution_cases(plan, "pending")] == ["Shop_TC004"]
    assert [case["id"] for case in P.execution_cases(plan, "retest")] == ["Shop_TC004"]
    with pytest.raises(ValueError):
        P.execution_cases(plan, "everything")


def test_workbook_marks(tmp_path):
    plan = build(tmp_path / "w.xlsx")
    decisions = {"Shop_TC001": "partial", "Shop_TC002": None, "Shop_TC004": "need_info", "Shop_TC005": "manual"}
    for case in plan["test_cases"]:
        case["automation"] = decisions[case["id"]]
    assert P.workbook_marks(plan) == {"Shop_TC001": "Automated", "Shop_TC004": "Not Automated",
                                      "Shop_TC005": "Not Automated"}


# ---------- writing the column ----------

def test_write_marks_touches_only_the_automation_column(tmp_path):
    source = build_workbook(tmp_path / "w.xlsx")
    out = tmp_path / "out.xlsx"
    result = write_automation_marks(source, out, "001-SHP_Shop", {
        "Shop_TC001": "Automated", "Shop_TC002": "Not Automated", "Shop_TC009": "Automated"})
    assert result == {"written": ["Shop_TC002"], "unchanged": ["Shop_TC001"], "missing": ["Shop_TC009"]}

    before = load_workbook(source)["001-SHP_Shop"]
    after = load_workbook(out)["001-SHP_Shop"]
    assert after[f"{AUTOMATION_COLUMN}10"].value == "Not Automated"
    for row in after.iter_rows(min_row=1, max_row=after.max_row):
        for cell in row:
            if cell.column_letter != AUTOMATION_COLUMN or cell.row < 9:
                assert cell.value == before[cell.coordinate].value, cell.coordinate
    assert after["L9"].value.startswith("=IF($A9")
    validations = [v for v in after.data_validations.dataValidation if f"{AUTOMATION_COLUMN}9" in v.sqref]
    assert validations and validations[0].formula1 == '"Automated,Not Automated"'


def test_write_marks_keeps_an_existing_drop_down(tmp_path):
    first = tmp_path / "first.xlsx"
    write_automation_marks(build_workbook(tmp_path / "w.xlsx"), first, "001-SHP_Shop", {})
    second = tmp_path / "second.xlsx"
    write_automation_marks(first, second, "001-SHP_Shop", {})
    sheet = load_workbook(second)["001-SHP_Shop"]
    assert sum(1 for v in sheet.data_validations.dataValidation if f"{AUTOMATION_COLUMN}9" in v.sqref) == 1


def test_write_marks_needs_the_column(tmp_path):
    source = build_workbook(tmp_path / "old.xlsx", automation_column=False)
    with pytest.raises(ValueError, match="no Automation column"):
        write_automation_marks(source, tmp_path / "out.xlsx", "001-SHP_Shop", {"Shop_TC001": "Automated"})


def test_write_marks_rejects_other_values(tmp_path):
    with pytest.raises(ValueError, match="not one of"):
        write_automation_marks(build_workbook(tmp_path / "w.xlsx"), tmp_path / "o.xlsx", "001-SHP_Shop",
                               {"Shop_TC001": "auto"})


# ---------- CLI ----------

def run_cli(*args, cwd):
    return subprocess.run([sys.executable, str(CLI), *args], cwd=cwd, capture_output=True, text=True, timeout=60)


def approve_all(plan_path, cwd, decisions):
    data = yaml.safe_load(plan_path.read_text(encoding="utf-8"))
    for case in data["test_cases"]:
        automation, review = decisions.get(case["id"], (case["automation"] or "auto", case["review"]))
        case["automation"], case["review"] = automation, review
    plan_path.write_text(yaml.safe_dump(data, allow_unicode=True, sort_keys=False), encoding="utf-8")
    assert run_cli("check", str(plan_path), cwd=cwd).returncode == 0
    approved = run_cli_tty("approve", str(plan_path), "--by", "NCM", cwd=cwd, answer="Shop")
    assert approved.returncode == 0, approved.output


@needs_pty
def test_cli_import_write_select(tmp_path):
    work = tmp_path / "work"
    work.mkdir()
    workbook = build_workbook(work / "cases.xlsx")
    imported = json.loads(run_cli("import", str(workbook), cwd=work).stdout)[0]
    assert imported["automation_from_workbook"] == ["Shop_TC001", "Shop_TC004"]
    plan_path = work / imported["plan"]

    for command in (("write-automation", str(plan_path), str(workbook)), ("select", str(plan_path))):
        refused = run_cli(*command, cwd=work)
        assert refused.returncode == 3 and "chưa được duyệt" in refused.stderr

    approve_all(plan_path, work, {"Shop_TC002": ("partial", "Toast text checked; layout by screenshot."),
                                  "Shop_TC005": ("need_info", "No steps.")})

    written = run_cli("write-automation", str(plan_path), str(workbook), cwd=work)
    assert written.returncode == 0, written.stderr
    result = json.loads(written.stdout)
    assert result["workbook"].endswith("cases_automation.xlsx")
    assert result["written"] == ["Shop_TC002", "Shop_TC005"]
    assert (result["automated"], result["not_automated"]) == (2, 2)
    out = load_workbook(work / "cases_automation.xlsx")["001-SHP_Shop"]
    assert [out[f"{AUTOMATION_COLUMN}{row}"].value for row in range(9, 13)] == [
        "Automated", "Automated", "Not Automated", "Not Automated"]
    assert load_workbook(workbook)["001-SHP_Shop"][f"{AUTOMATION_COLUMN}10"].value is None, "input left as it was"

    assert run_cli("select", str(plan_path), cwd=work).stdout.split() == ["Shop_TC001", "Shop_TC002"]
    assert run_cli("select", str(plan_path), "--format", "pytest-k", cwd=work).stdout.strip() == "Shop_TC001 or Shop_TC002"
    rows = json.loads(run_cli("select", str(plan_path), "--format", "json", cwd=work).stdout)
    assert [(row["id"], row["automation"], row["state"]) for row in rows] == [
        ("Shop_TC001", "auto", "passed"), ("Shop_TC002", "partial", "passed")]
    empty = run_cli("select", str(plan_path), "--scope", "retest", cwd=work)
    assert empty.returncode == 0 and empty.stdout.strip() == "" and "Không có case Automated" in empty.stderr

    # Importing the written workbook changes nothing: the ticks now match the approved decisions.
    again = json.loads(run_cli("import", str(work / "cases_automation.xlsx"), cwd=work).stdout)[0]
    assert again["automation_from_workbook"] == [] and again["changed_since_last_import"] == []
    assert again["approval"]["state"] == "approved"
    assert again["summary"]["by_workbook_automation"] == {"Automated": 2, "Not Automated": 2}
    assert again["summary"]["issues"]["warning"] == imported["summary"]["issues"]["warning"]


@needs_pty
def test_cli_write_refuses_a_changed_workbook(tmp_path):
    work = tmp_path / "work"
    work.mkdir()
    workbook = build_workbook(work / "cases.xlsx")
    plan_path = work / json.loads(run_cli("import", str(workbook), cwd=work).stdout)[0]["plan"]
    approve_all(plan_path, work, {"Shop_TC005": ("need_info", "No steps.")})

    edited = list(CASES)
    edited[0] = edited[0][:2] + ("List shows all products",) + edited[0][3:]
    build_workbook(workbook, edited)
    refused = run_cli("write-automation", str(plan_path), str(workbook), cwd=work)
    assert refused.returncode == 3 and "import lại" in refused.stderr


@needs_pty
def test_cli_write_in_place(tmp_path):
    work = tmp_path / "work"
    work.mkdir()
    workbook = build_workbook(work / "cases.xlsx")
    plan_path = work / json.loads(run_cli("import", str(workbook), cwd=work).stdout)[0]["plan"]
    approve_all(plan_path, work, {"Shop_TC005": ("need_info", "No steps.")})
    before = hashlib.sha256(workbook.read_bytes()).hexdigest()
    done = run_cli("write-automation", str(plan_path), str(workbook), "--in-place", cwd=work)
    assert done.returncode == 0, done.stderr
    assert hashlib.sha256(workbook.read_bytes()).hexdigest() != before
    assert not (work / "cases_automation.xlsx").exists()
    assert load_workbook(workbook)["001-SHP_Shop"][f"{AUTOMATION_COLUMN}12"].value == "Not Automated"
