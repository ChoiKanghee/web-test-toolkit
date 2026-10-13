"""Read test cases from the team test case workbook (結合テスト仕様書兼結果書 template).

The template is located by its bilingual headers, not by fixed cells: a column header
ends with an English label in parentheses, e.g. "テストケース番号 *\n(TC ID)", or is the
English label itself ("Automation").

This module writes only two things, each after a tester decision: the Automation column
(write_automation_marks) and the content of accepted AI proposals (write_case_changes).
"""

import math
import re
import unicodedata
from copy import copy
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet


HEADER_SCAN_ROWS = 20

# English label inside the header parentheses -> field name
CASE_COLUMNS = {
    "tc id": "id",
    "category": "category",
    "title": "title",
    "pre-condition": "precondition",
    "precondition": "precondition",
    "steps": "steps",
    "test data": "test_data",
    "expected result": "expected",
    "priority": "priority",
    "final result": "final_result",
    "tester / pic": "tester",
    "tester": "tester",
    "date": "date",
    "bug id": "bug_id",
    "note": "note",
    "automation": "automation_mark",
}
RUN_LABELS = ("1st", "2nd", "3rd", "4th", "5th")

META_LABELS = {
    "screen / function name": "function_name",
    "module": "module",
    "environment": "environment",
}

DATA_COLUMNS = {
    "data id": "id",
    "type": "type",
    "content": "content",
    "language": "language",
    "used in": "used_in",
    "note": "note",
}

_PAREN = re.compile(r"\(([^()]*)\)\s*$")


def english_label(value) -> Optional[str]:
    """'テストケース番号 *\\n(TC ID)' -> 'tc id'; 'テスト結果 (Test Result) – Chrome' -> 'test result'."""
    if not isinstance(value, str):
        return None
    text = value.strip()
    match = _PAREN.search(text)
    if not match:
        match = re.search(r"\(([^()]*)\)", text)
    return match.group(1).strip().lower() if match else None


def case_column(value) -> Optional[str]:
    """Field name of a test case column header, or None.

    Most headers carry the English label in parentheses; a header written in English only
    ("Automation") is matched as is.
    """
    label = english_label(value)
    if label is None and isinstance(value, str) and "(" not in value:
        label = " ".join(value.split()).lower()
    return CASE_COLUMNS.get(label or "")


def clean(value) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).replace("\r\n", "\n").strip()
    return text or None


@dataclass
class RawCase:
    row: int
    values: dict
    runs: list = field(default_factory=list)


@dataclass
class RawSheet:
    name: str
    meta: dict
    columns: dict
    cases: list


@dataclass
class RawDataItem:
    row: int
    values: dict


def _find_header_row(ws: Worksheet, wanted: str) -> Optional[int]:
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, HEADER_SCAN_ROWS)):
        for cell in row:
            if english_label(cell.value) == wanted:
                return cell.row
    return None


def _find_case_header_row(ws: Worksheet) -> Optional[int]:
    """The row holding the TC ID, Steps and Expected Result headers together.

    Other sheets may mention '(TC ID)' (e.g. the Item Description sheet) without being a case sheet.
    """
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, HEADER_SCAN_ROWS)):
        labels = {case_column(cell.value) for cell in row}
        if {"id", "steps", "expected"} <= labels:
            return row[0].row
    return None


def is_case_sheet(ws: Worksheet) -> bool:
    return _find_case_header_row(ws) is not None


def _sheet_meta(ws: Worksheet, header_row: int) -> dict:
    meta = {}
    for row in ws.iter_rows(min_row=1, max_row=header_row - 1):
        cells = [cell for cell in row if cell.value not in (None, "")]
        for index, cell in enumerate(cells):
            name = META_LABELS.get(english_label(cell.value) or "")
            if name and index + 1 < len(cells):
                value = clean(cells[index + 1].value)
                if value and META_LABELS.get(english_label(value) or "") is None:
                    meta[name] = value
    return meta


@dataclass
class CaseLayout:
    header_row: int
    first_data_row: int
    columns: dict
    run_columns: list


def case_layout(ws: Worksheet) -> CaseLayout:
    header_row = _find_case_header_row(ws)
    if header_row is None:
        raise ValueError(f"Sheet '{ws.title}' is not a test case sheet (no row with TC ID, Steps and Expected Result headers).")

    columns = {}
    for cell in ws[header_row]:
        name = case_column(cell.value)
        if name and name not in columns:
            columns[name] = cell.column

    run_columns = []
    sub_row = header_row + 1
    for cell in ws[sub_row]:
        label = english_label(cell.value)
        if label in RUN_LABELS:
            run_columns.append((RUN_LABELS.index(label), cell.column))
    run_columns = [column for _, column in sorted(run_columns)]
    first_data_row = sub_row + 1 if run_columns else header_row + 1
    return CaseLayout(header_row=header_row, first_data_row=first_data_row, columns=columns, run_columns=run_columns)


def read_case_sheet(ws: Worksheet) -> RawSheet:
    layout = case_layout(ws)
    header_row, first_data_row = layout.header_row, layout.first_data_row
    columns, run_columns = layout.columns, layout.run_columns

    missing = [name for name in ("id", "title", "steps", "expected") if name not in columns]
    if missing:
        raise ValueError(f"Sheet '{ws.title}': missing columns {', '.join(missing)}.")

    cases = []
    for row in range(first_data_row, ws.max_row + 1):
        case_id = clean(ws.cell(row=row, column=columns["id"]).value)
        if not case_id:
            continue
        values = {name: clean(ws.cell(row=row, column=column).value) for name, column in columns.items()}
        runs = [clean(ws.cell(row=row, column=column).value) for column in run_columns]
        cases.append(RawCase(row=row, values=values, runs=runs))

    return RawSheet(name=ws.title, meta=_sheet_meta(ws, header_row), columns=columns, cases=cases)


def read_test_data(ws: Worksheet) -> list:
    header_row = _find_header_row(ws, "data id")
    if header_row is None:
        return []

    columns = {}
    for cell in ws[header_row]:
        name = DATA_COLUMNS.get(english_label(cell.value) or "")
        if name and name not in columns:
            columns[name] = cell.column

    items = []
    for row in range(header_row + 1, ws.max_row + 1):
        data_id = clean(ws.cell(row=row, column=columns["id"]).value)
        if not data_id:
            continue
        values = {name: clean(ws.cell(row=row, column=column).value) for name, column in columns.items()}
        items.append(RawDataItem(row=row, values=values))
    return items


@dataclass
class RawWorkbook:
    path: Path
    case_sheets: list
    test_data: list


def read_workbook(path: Path, sheets: Optional[list] = None) -> RawWorkbook:
    workbook = load_workbook(path, data_only=False)

    if sheets:
        unknown = [name for name in sheets if name not in workbook.sheetnames]
        if unknown:
            raise ValueError(f"Sheet not found: {', '.join(unknown)}. Sheets: {', '.join(workbook.sheetnames)}")
        case_sheets = [read_case_sheet(workbook[name]) for name in sheets]
    else:
        case_sheets = [read_case_sheet(ws) for ws in workbook.worksheets if is_case_sheet(ws)]
        if not case_sheets:
            raise ValueError(f"No test case sheet (with a '(TC ID)' header) found in {path.name}.")

    test_data = []
    for ws in workbook.worksheets:
        if _find_header_row(ws, "data id") is not None and not is_case_sheet(ws):
            test_data.extend(read_test_data(ws))

    return RawWorkbook(path=Path(path), case_sheets=case_sheets, test_data=test_data)


# ========= Write: the Automation column only =========

AUTOMATION_CHOICES = ("Automated", "Not Automated")


def _has_validation(ws: Worksheet, cell_ref: str) -> bool:
    return any(cell_ref in validation.sqref for validation in ws.data_validations.dataValidation)


def write_automation_marks(path: Path, out_path: Path, sheet_name: str, marks: dict) -> dict:
    """Write Automated / Not Automated into the sheet's Automation column, by TC ID.

    marks: {test_id: "Automated" | "Not Automated"}. Every other cell is left as it is, formulas
    included. When the column has no drop-down yet, one is added with the two values.
    Returns {"written": [...], "unchanged": [...], "missing": [...]} (lists of test ids).
    """
    workbook = load_workbook(path, data_only=False)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Sheet not found: {sheet_name}. Sheets: {', '.join(workbook.sheetnames)}")
    ws = workbook[sheet_name]
    layout = case_layout(ws)
    if "automation_mark" not in layout.columns:
        raise ValueError(f"Sheet '{sheet_name}' has no Automation column (header 'Automation' in the TC ID header row).")
    id_column = layout.columns["id"]
    mark_column = layout.columns["automation_mark"]

    rows = {}
    for row in range(layout.first_data_row, ws.max_row + 1):
        case_id = clean(ws.cell(row=row, column=id_column).value)
        if case_id and case_id not in rows:
            rows[case_id] = row

    result = {"written": [], "unchanged": [], "missing": []}
    for case_id, mark in marks.items():
        if mark not in AUTOMATION_CHOICES:
            raise ValueError(f"{case_id}: '{mark}' is not one of {', '.join(AUTOMATION_CHOICES)}.")
        row = rows.get(case_id)
        if row is None:
            result["missing"].append(case_id)
            continue
        cell = ws.cell(row=row, column=mark_column)
        if clean(cell.value) == mark:
            result["unchanged"].append(case_id)
        else:
            cell.value = mark
            result["written"].append(case_id)

    letter = ws.cell(row=layout.first_data_row, column=mark_column).column_letter
    if not _has_validation(ws, f"{letter}{layout.first_data_row}"):
        last_row = max([layout.first_data_row, ws.max_row] + [
            cell_range.max_row for validation in ws.data_validations.dataValidation for cell_range in validation.sqref.ranges])
        validation = DataValidation(type="list", formula1='"' + ",".join(AUTOMATION_CHOICES) + '"', allow_blank=True)
        validation.add(f"{letter}{layout.first_data_row}:{letter}{last_row}")
        ws.add_data_validation(validation)

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    workbook.save(out_path)
    return result


# ========= Write: accepted AI proposals =========

# Plan field -> workbook column field
PLAN_TO_COLUMN = {"title": "title", "category": "category", "preconditions": "precondition", "steps": "steps",
                  "test_data": "test_data", "expected": "expected", "priority": "priority"}
ALWAYS_NUMBERED = ("steps", "expected")
TEMPLATE_PRIORITY = {"high": "High", "medium": "Medium", "low": "Low"}
LINE_HEIGHT = 15.0


def cell_text(column_field: str, value):
    """Format a value the way the template writes it: steps and expected results as '1. ...' lines,
    a precondition numbered only when it has several items, priority as High / Medium / Low."""
    if value is None:
        return None
    if isinstance(value, list):
        items = [str(item).strip() for item in value if str(item).strip()]
        if not items:
            return None
        if len(items) == 1 and column_field not in ALWAYS_NUMBERED:
            return items[0]
        return "\n".join(f"{index}. {item}" for index, item in enumerate(items, start=1))
    text = str(value).strip()
    if column_field == "priority":
        return TEMPLATE_PRIORITY.get(text.lower(), text)
    return text or None


def _display_width(text: str) -> int:
    return sum(2 if unicodedata.east_asian_width(char) in ("W", "F") else 1 for char in text)


def _row_height(ws: Worksheet, row: int, columns: list) -> float:
    lines = 1
    for column in columns:
        value = ws.cell(row=row, column=column).value
        if not isinstance(value, str) or value.startswith("="):
            continue
        width = ws.column_dimensions[ws.cell(row=row, column=column).column_letter].width or 8.43
        per_line = max(1, int(width * 1.1))
        lines = max(lines, sum(max(1, math.ceil(_display_width(part) / per_line)) for part in value.split("\n")))
    return LINE_HEIGHT * lines + 3


def _append_note(cell, line: str) -> None:
    current = clean(cell.value)
    cell.value = f"{current}\n{line}" if current else line


def write_case_changes(path: Path, out_path: Path, sheet_name: str, edits: list, new_cases: list, notes: dict) -> dict:
    """Write accepted proposals into a test case sheet.

    edits:     [(test_id, plan_field, value)]   value: str, or list for preconditions / steps / expected
    new_cases: [{"id", "category", "title", "precondition", "steps", "test_data", "expected", "priority",
                 "automation_mark"}]            appended after the last case, in the template's formatted rows
    notes:     {test_id: line}                  appended to the Note column, so the team sees what came from AI
    Nothing is written when an edited case is missing or a new id already exists.
    """
    workbook = load_workbook(path, data_only=False)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Sheet not found: {sheet_name}. Sheets: {', '.join(workbook.sheetnames)}")
    ws = workbook[sheet_name]
    layout = case_layout(ws)
    columns = layout.columns

    rows = {}
    last_case_row = None
    for row in range(layout.first_data_row, ws.max_row + 1):
        case_id = clean(ws.cell(row=row, column=columns["id"]).value)
        if case_id:
            rows.setdefault(case_id, row)
            last_case_row = row

    problems = []
    for test_id, plan_field, _ in edits:
        if test_id not in rows:
            problems.append(f"{test_id}: không có trong sheet.")
        if plan_field not in PLAN_TO_COLUMN:
            problems.append(f"{test_id}: không ghi được trường '{plan_field}'.")
        elif PLAN_TO_COLUMN[plan_field] not in columns:
            problems.append(f"{test_id}: sheet không có cột cho '{plan_field}'.")
    for case in new_cases:
        if case["id"] in rows:
            problems.append(f"{case['id']}: mã đã có trong sheet.")
        if case.get("automation_mark") and "automation_mark" not in columns:
            problems.append(f"{case['id']}: sheet không có cột Automation.")
    if problems:
        raise ValueError("Không ghi được vào workbook:\n  - " + "\n  - ".join(problems))

    result = {"edited": [], "added": [], "beyond_template_rows": []}
    for test_id, plan_field, value in edits:
        column_field = PLAN_TO_COLUMN[plan_field]
        ws.cell(row=rows[test_id], column=columns[column_field]).value = cell_text(column_field, value)
        result["edited"].append([test_id, plan_field])

    style_row = last_case_row or layout.first_data_row
    template_last = max([style_row] + [cell_range.max_row for validation in ws.data_validations.dataValidation
                                       for cell_range in validation.sqref.ranges])
    last_column = max(cell.column for cell in ws[layout.header_row] if cell.value is not None)
    last_column = max([last_column] + list(columns.values()) + layout.run_columns)
    next_row = (last_case_row or layout.first_data_row - 1) + 1
    for case in new_cases:
        while clean(ws.cell(row=next_row, column=columns["id"]).value):
            next_row += 1
        row = next_row
        for column in range(1, last_column + 1):
            source = ws.cell(row=style_row, column=column)
            target = ws.cell(row=row, column=column)
            if source.has_style:
                target.font, target.fill, target.border = copy(source.font), copy(source.fill), copy(source.border)
                target.alignment, target.number_format = copy(source.alignment), source.number_format
                target.protection = copy(source.protection)
            if target.value is None and isinstance(source.value, str) and source.value.startswith("="):
                target.value = Translator(source.value, origin=source.coordinate).translate_formula(target.coordinate)
        for column_field, value in case.items():
            if column_field in columns and value not in (None, [], ""):
                ws.cell(row=row, column=columns[column_field]).value = cell_text(column_field, value)
        ws.row_dimensions[row].height = _row_height(ws, row, list(columns.values()))
        result["added"].append([case["id"], row])
        if row > template_last:
            result["beyond_template_rows"].append(case["id"])
        next_row = row + 1

    for test_id, row in list(rows.items()) + [(case_id, row) for case_id, row in result["added"]]:
        if test_id in notes and "note" in columns:
            _append_note(ws.cell(row=row, column=columns["note"]), notes[test_id])

    for test_id in {test_id for test_id, _ in result["edited"]}:
        row = rows[test_id]
        needed = _row_height(ws, row, list(columns.values()))
        current = ws.row_dimensions[row].height
        if current is None or current < needed:
            ws.row_dimensions[row].height = needed

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    workbook.save(out_path)
    return result
